import pandas as pd
from langchain.vectorstores import FAISS
from langchain.embeddings import HuggingFaceEmbeddings
from langchain.schema import Document
from langchain.prompts import PromptTemplate
from langchain.chains import RetrievalQA
from langchain_ollama import OllamaLLM
from langchain.evaluation import load_evaluator
from openpyxl import load_workbook
import os

def load_qa_excel(file_path):
    df = pd.read_excel(file_path)
    question_titles = df['Question Title'].tolist()
    question_bodies = df['Question Body'].tolist()
    accepted_answers = df['Accepted Answer Body'].tolist()
    return question_titles, question_bodies, accepted_answers

def load_query_excel(file_path):
    df = pd.read_excel(file_path)
    query_titles = df['Question Title'].tolist()
    query_bodies = df['Question Body'].tolist()
    queries = [f"Title: {title}\nBody: {body}" for title, body in zip(query_titles, query_bodies)]
    return df, queries


def create_documents(question_titles, question_bodies, accepted_answers):
    documents = []
    for i, (title, body, answer) in enumerate(zip(question_titles, question_bodies, accepted_answers), 1):
        text = f"Question Title: {title}\nQuestion Body: {body}\nAccepted Answer: {answer}"
        documents.append(Document(page_content=text))
        print(f"Processed document {i}/{len(question_titles)}")
    return documents

def create_embeddings(documents):
    print("Setting up the FAISS vector store with embeddings...")
    embedder = HuggingFaceEmbeddings()
    vector_store = FAISS.from_documents(documents, embedder)
    print("FAISS vector store setup completed.")
    return vector_store, embedder


def get_context_if_relevant(query, vector_store, embedder, evaluator):
    retrieved_docs = vector_store.similarity_search(query, k=1)
    if retrieved_docs:
        context = retrieved_docs[0].page_content
        distance_result = evaluator.evaluate_strings(prediction=query, reference=context)
        distance = distance_result['score']
        return context, distance
    return None, 0.0

def create_rag_chain(model_name, vector_store):
    template = """
    You are a helpful assistant with expertise in coding and technical topics. Use the following context to answer the question as accurately as possible, especially focusing on technical details and code examples if relevant.
    
    Context: {context}
    
    Question: {question}
    
    Provide a helpful and accurate answer, focusing on coding and technical topics.
    """
    prompt_template = PromptTemplate(input_variables=["context", "question"], template=template)
    llm = OllamaLLM(model=model_name)
    rag_chain = RetrievalQA.from_chain_type(
        llm=llm,
        chain_type="stuff",
        retriever=vector_store.as_retriever(),
        return_source_documents=True,
        chain_type_kwargs={'prompt': prompt_template}
    )
    return rag_chain

def append_to_excel(df, output_file):
    if os.path.exists(output_file):
        with pd.ExcelWriter(output_file, mode='a', if_sheet_exists='overlay', engine='openpyxl') as writer:
            startrow = writer.sheets['Sheet1'].max_row
            df.to_excel(writer, index=False, header=False, startrow=startrow)
    else:
        df.to_excel(output_file, index=False)

def main(qa_file, query_file, output_file):
    question_titles, question_bodies, accepted_answers = load_qa_excel(qa_file)
    query_df, queries = load_query_excel(query_file)

    documents = create_documents(question_titles, question_bodies, accepted_answers)

    vector_store, embedder = create_embeddings(documents)
    hf_evaluator = load_evaluator("embedding_distance", embeddings=embedder)

    rag_chain_codellama = create_rag_chain("codellama:latest", vector_store)
    rag_chain_solar = create_rag_chain("solar:latest", vector_store)
    rag_chain_mistral = create_rag_chain("mistral:latest", vector_store)

    for i, query in enumerate(queries, 1):
        context, distance = get_context_if_relevant(query, vector_store, embedder, hf_evaluator)
        query_context = context if context and distance < 0.5 else ""

        response_codellama = rag_chain_codellama({"query": query, "context": query_context})
        response_solar = rag_chain_solar({"query": query, "context": query_context})
        response_mistral = rag_chain_mistral({"query": query, "context": query_context})

        query_df.loc[i - 1, 'codellama_with_rag'] = response_codellama['result']
        query_df.loc[i - 1, 'solar_with_rag'] = response_solar['result']
        query_df.loc[i - 1, 'mistral_with_rag'] = response_mistral['result']

        append_to_excel(query_df.iloc[[i - 1]], output_file)
        print(f"Processed and saved query {i}/{len(queries)}")


if __name__ == "__main__":
    qa_file = '../Dataset/DataExcel/newStackExchangeQsAnswer/stackExchangeQsAndAnswersDB.xlsx'
    query_file = '../Dataset/DataExcel/newStackExchangeQsAnswer/stackExchangeQsAndAnswersTest.xlsx'
    output_file = '../Dataset/updated_stackexchange_results.xlsx'
    
    main(qa_file, query_file, output_file)